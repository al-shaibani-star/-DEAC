# -*- coding: utf-8 -*-
"""
Re-skin paper/DEAC_paper.docx into the IJIES template format (A4, two-column body
with a single-column title block, Times New Roman, template front matter, numbered
bold section headings). Preserves all equations/tables/figures from the source.
Output: paper/DEAC_paper_IJIES.docx
"""
import copy
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "paper/DEAC_paper.docx"
OUT = "paper/DEAC_paper_IJIES.docx"

TITLE = "DEAC: A Defensive Ensemble Framework for Automatic Graph-Based Clustering of High-Dimensional Data"
AUTHOR = "Hussein Ali Ismael¹*"
AFFIL = "¹Department of Computer Science, College of Science, University of Diyala, Diyala, Iraq"
EMAIL = "* Corresponding author’s Email: Al-Shaibani@uodiyala.edu.iq"
ABSTRACT = ("Clustering high-dimensional data without a known cluster count conventionally returns the "
    "partition that maximizes an internal quality score such as modularity or silhouette. We show this is "
    "unsafe: across twelve benchmarks, metaheuristically optimizing internal objectives degrades external "
    "agreement, because internal and external quality are only weakly correlated. We study this misalignment "
    "directly. Using a label-aware oracle that selects the best engine per dataset as an upper bound, we "
    "measure how much any label-free rule recovers: a three-layer QualityGate recovers about two-thirds of "
    "the oracle’s gain, and learned, stability, and consensus selectors never close the gap, indicating it "
    "is structural. The multi-representation Shared Nearest Neighbor and Leiden ensemble the gate arbitrates "
    "is itself a competitive label-free method that infers the cluster count, remaining statistically "
    "indistinguishable from strong classical and deep baselines while trailing only one recent method. We "
    "package these as DEAC and release all per-seed results and scripts.")
KEYWORDS = ("Clustering, High-dimensional data, Graph-based clustering, Differential evolution, "
    "Shared nearest neighbors, Leiden algorithm, Ensemble clustering")


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = "Times New Roman"
    r = run._element.rPr
    if r is None:
        r = run._element.get_or_add_rPr()
    rf = r.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); r.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)  # force black (remove pandoc theme color)


def hrule(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom); pPr.append(pbdr)


def make_sectPr(num_cols, continuous):
    s = OxmlElement("w:sectPr")
    if continuous:
        t = OxmlElement("w:type"); t.set(qn("w:val"), "continuous"); s.append(t)
    pgsz = OxmlElement("w:pgSz"); pgsz.set(qn("w:w"), "11906"); pgsz.set(qn("w:h"), "16838"); pgsz.set(qn("w:code"), "9")
    s.append(pgsz)
    pgmar = OxmlElement("w:pgMar")
    for k, v in [("top", "1440"), ("right", "1080"), ("bottom", "1440"), ("left", "1080"),
                 ("header", "850"), ("footer", "567"), ("gutter", "0")]:
        pgmar.set(qn("w:" + k), v)
    s.append(pgmar)
    cols = OxmlElement("w:cols"); cols.set(qn("w:num"), str(num_cols)); cols.set(qn("w:space"), "432")
    s.append(cols)
    return s


def main():
    doc = Document(SRC)

    # --- defaults: Times New Roman 10pt body ---
    for sname, size in [("Normal", 10), ("Body Text", 10), ("First Paragraph", 10)]:
        try:
            st = doc.styles[sname]; st.font.name = "Times New Roman"; st.font.size = Pt(size)
            st.element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        except Exception:
            pass

    # black-out colored cross-reference hyperlinks (pandoc colors them)
    try:
        hl = doc.styles["Hyperlink"]; hl.font.color.rgb = RGBColor(0, 0, 0); hl.font.underline = False
    except Exception:
        pass
    for hy in doc.element.body.iter(qn("w:hyperlink")):
        for rpr in hy.iter(qn("w:rPr")):
            for col in rpr.findall(qn("w:color")):
                col.set(qn("w:val"), "000000")

    paras = doc.paragraphs
    intro = next(p for p in paras if p.text.strip().lower() == "introduction")

    # --- build template front matter above Introduction ---
    def add(text, size, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, rule=False):
        p = intro.insert_paragraph_before("")
        p.alignment = align
        p.paragraph_format.space_after = Pt(2)
        if text:
            run = p.add_run(text); set_font(run, size, bold, italic)
        if rule:
            hrule(p)
        return p

    add(TITLE, 14, bold=True)
    add(AUTHOR, 11, bold=True)
    add(AFFIL, 11, italic=True)
    email_p = add(EMAIL, 10)
    hrule(email_p)                 # line below author info (above abstract)
    ab = add("", 10, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    r = ab.add_run("Abstract: "); set_font(r, 10, bold=True)
    r = ab.add_run(ABSTRACT); set_font(r, 10)
    kw = add("", 10, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = kw.add_run("Keywords: "); set_font(r, 10, bold=True)
    r = kw.add_run(KEYWORDS); set_font(r, 10)
    hrule(kw)                      # line below keywords closes the title block
    rule_p = kw                    # the 1-col section ends at the keywords paragraph

    # --- delete OLD front matter (original paras 0-4: Title..Keywords) ---
    for p in list(doc.paragraphs):
        if p._p is intro._p:
            break
        if p.style.name in ("Title", "Author", "Abstract Title", "Abstract") or \
           (p.style.name == "First Paragraph" and p.text.startswith("Clustering, High-dimensional")):
            p._p.getparent().remove(p._p)

    # --- number + restyle section headings ---
    h1 = 0
    for p in doc.paragraphs:
        sn = p.style.name
        if sn == "Heading 1":
            h1 += 1
            txt = p.text
            p.clear(); run = p.add_run(f"{h1}.  {txt}"); set_font(run, 12, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif sn == "Heading 2":
            txt = p.text
            p.clear(); run = p.add_run(txt); set_font(run, 11, bold=True)
        elif sn == "Heading 3":
            txt = p.text
            p.clear(); run = p.add_run(txt); set_font(run, 10, bold=True, italic=True)

    # --- sections: title block single-col, body two-col ---
    # final (doc) sectPr -> 2 columns
    body_sect = doc.element.body.find(qn("w:sectPr"))
    if body_sect is not None:
        body_sect.getparent().remove(body_sect)
    doc.element.body.append(make_sectPr(2, continuous=True))  # body continues on page 1 after title block
    # put a 1-col continuous sectPr in the closing-rule paragraph (ends the title-block section)
    pPr = rule_p._p.get_or_add_pPr()
    pPr.append(make_sectPr(1, continuous=True))

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
